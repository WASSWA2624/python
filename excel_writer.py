"""Workbook create / upsert / format (spec 10).

The rules that shape this module:

* **Exactly one workbook per match date**, named ``YYYY-MM-DD.xlsx`` from the
  match date alone. The window, the thresholds and the run time never affect
  the filename, and no suffixed variant is ever produced - not even when the
  target is locked.
* **Rows are upserted on ``match_key``** and never silently deleted. A fixture
  that disappears from the sportsbook, or falls outside a narrowed window,
  stays visible with its history intact.
* **Writes are atomic**: the workbook is built beside the target and moved onto
  it with ``os.replace``, so a crash mid-write cannot leave a corrupt file.
* **Sheets this module does not own are preserved**, so notes added by hand
  survive a re-run.

openpyxl does all reading and writing; pandas is not needed here and is used
for computation only elsewhere.
"""

from __future__ import annotations

import os
import time as time_module
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from config import DISCLAIMER, Config
from models import (
    STATUS_NOT_IN_LATEST_RUN,
    UTC,
    AnalysisWindow,
    MatchRecord,
    RunLogEntry,
    SourceRecord,
    now_utc,
    to_eat,
    to_utc,
)
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from ranking import ranked_selections
from utils.logging import get_logger

log = get_logger("excel")

SHEET_SUMMARY = "Summary"
SHEET_RANKED = "Ranked Selections"
SHEET_ALL = "All Matches"
SHEET_RUN_LOG = "Run Log"
SHEET_SOURCES = "Data Sources"

#: Sheets this module owns and rewrites. Anything else in the file is left alone.
OWNED_SHEETS = (SHEET_SUMMARY, SHEET_RANKED, SHEET_ALL, SHEET_RUN_LOG, SHEET_SOURCES)

FORMAT_PERCENT = "0.0%"
FORMAT_ODDS = "0.00"
FORMAT_GOALS = "0.00"
FORMAT_SCORE = "0.000"
FORMAT_SIGNED_PERCENT = "+0.0%;-0.0%;0.0%"
FORMAT_DATETIME = "yyyy-mm-dd hh:mm"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True)


class WorkbookLockedError(RuntimeError):
    """The target workbook is open elsewhere and cannot be replaced."""


# ==========================================================================
# Column definitions
# ==========================================================================
@dataclass(frozen=True)
class Column:
    """One workbook column and how a record maps onto it."""

    header: str
    attribute: str
    width: int = 14
    number_format: str | None = None
    to_cell: Callable[[Any], Any] | None = None
    from_cell: Callable[[Any], Any] | None = None


def _eat_naive(value: datetime | None) -> Any:
    """Excel cannot hold a tz-aware datetime; render EAT wall-clock time."""
    return None if value is None else to_eat(value).replace(tzinfo=None)


def _utc_naive(value: datetime | None) -> Any:
    return None if value is None else to_utc(value).replace(tzinfo=None)


def _read_utc(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=UTC) if value.tzinfo is None else to_utc(value)
    if isinstance(value, str) and value.strip():
        try:
            parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError:
            return None
        return parsed.replace(tzinfo=UTC) if parsed.tzinfo is None else to_utc(parsed)
    return None


def _read_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _read_int(value: Any) -> int | None:
    number = _read_float(value)
    return None if number is None else int(round(number))


def _read_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"true", "yes", "1"}


def _read_str(value: Any) -> str:
    return "" if value is None else str(value)


def _pct(number_format: str = FORMAT_PERCENT) -> dict[str, Any]:
    return {"number_format": number_format, "from_cell": _read_float}


#: ``All Matches`` - every fixture, including excluded ones, with the full set
#: of section 4 statistics and an exclusion reason (spec 10.4).
ALL_MATCHES_COLUMNS: tuple[Column, ...] = (
    Column("Match Key", "match_key", 34, from_cell=_read_str),
    Column("Match", "", 30, from_cell=_read_str),  # derived; not read back
    Column("League", "league", 24, from_cell=_read_str),
    Column("Home Team", "home_team", 20, from_cell=_read_str),
    Column("Away Team", "away_team", 20, from_cell=_read_str),
    Column("Kick-off (EAT)", "kickoff_utc", 18, FORMAT_DATETIME, to_cell=_eat_naive),
    Column(
        "Kick-off (UTC)",
        "kickoff_utc",
        18,
        FORMAT_DATETIME,
        to_cell=_utc_naive,
        from_cell=_read_utc,
    ),
    Column("Status", "status", 22, from_cell=_read_str),
    Column("Exclusion Reason", "exclusion_reason", 46, from_cell=_read_str),
    Column("Rank", "rank", 7, from_cell=_read_int),
    Column("Category", "category", 20, from_cell=_read_str),
    Column("Qualified", "qualified", 10, from_cell=_read_bool),
    Column("Score", "score", 9, FORMAT_SCORE, from_cell=_read_float),
    Column("O1.5 Odds", "odds_over_15", 10, FORMAT_ODDS, from_cell=_read_float),
    Column("U1.5 Odds", "odds_under_15", 10, FORMAT_ODDS, from_cell=_read_float),
    Column("O2.5 Odds", "odds_over_25", 10, FORMAT_ODDS, from_cell=_read_float),
    Column("U2.5 Odds", "odds_under_25", 10, FORMAT_ODDS, from_cell=_read_float),
    Column("1", "odds_home", 8, FORMAT_ODDS, from_cell=_read_float),
    Column("X", "odds_draw", 8, FORMAT_ODDS, from_cell=_read_float),
    Column("2", "odds_away", 8, FORMAT_ODDS, from_cell=_read_float),
    Column("BTTS Yes", "odds_btts_yes", 10, FORMAT_ODDS, from_cell=_read_float),
    Column("BTTS No", "odds_btts_no", 10, FORMAT_ODDS, from_cell=_read_float),
    Column("Model Prob", "model_probability", 12, **_pct()),
    Column("Lambda Home", "lambda_home", 12, FORMAT_GOALS, from_cell=_read_float),
    Column("Lambda Away", "lambda_away", 12, FORMAT_GOALS, from_cell=_read_float),
    Column("Lambda Total", "lambda_total", 12, FORMAT_GOALS, from_cell=_read_float),
    Column("Confidence", "confidence", 12, from_cell=_read_str),
    Column("Data Completeness", "data_completeness", 17, **_pct()),
    Column("Implied (Raw)", "implied_raw", 13, **_pct()),
    Column("Implied (Fair)", "implied_fair", 13, **_pct()),
    Column("Overround", "overround", 11, FORMAT_GOALS, from_cell=_read_float),
    Column("Margin Adjusted", "margin_adjusted", 15, from_cell=_read_bool),
    Column("Value", "value", 10, FORMAT_SIGNED_PERCENT, from_cell=_read_float),
    Column("EV", "ev", 10, FORMAT_SIGNED_PERCENT, from_cell=_read_float),
    Column("Kelly", "kelly", 9, FORMAT_PERCENT, from_cell=_read_float),
    Column("H2H Matches", "h2h_matches", 12, from_cell=_read_int),
    Column("H2H O1.5 %", "h2h_over15_rate", 12, **_pct()),
    Column("H2H Avg Goals", "h2h_avg_goals", 14, FORMAT_GOALS, from_cell=_read_float),
    Column("H2H 2+ Goals", "h2h_2plus", 13, from_cell=_read_int),
    Column("H2H Recent", "h2h_recent", 50, from_cell=_read_str),
    Column("Home Form Matches", "home_form_matches", 17, from_cell=_read_int),
    Column("Home Form O1.5 %", "home_form_over15_rate", 17, **_pct()),
    Column("Home Avg Scored", "home_avg_scored", 16, FORMAT_GOALS, from_cell=_read_float),
    Column("Home Avg Conceded", "home_avg_conceded", 17, FORMAT_GOALS, from_cell=_read_float),
    Column("Home Avg Total", "home_avg_total", 15, FORMAT_GOALS, from_cell=_read_float),
    Column("Home Form 2+", "home_form_2plus", 13, from_cell=_read_int),
    Column("Away Form Matches", "away_form_matches", 17, from_cell=_read_int),
    Column("Away Form O1.5 %", "away_form_over15_rate", 17, **_pct()),
    Column("Away Avg Scored", "away_avg_scored", 16, FORMAT_GOALS, from_cell=_read_float),
    Column("Away Avg Conceded", "away_avg_conceded", 17, FORMAT_GOALS, from_cell=_read_float),
    Column("Away Avg Total", "away_avg_total", 15, FORMAT_GOALS, from_cell=_read_float),
    Column("Away Form 2+", "away_form_2plus", 13, from_cell=_read_int),
    Column("Home (Home) Matches", "home_home_matches", 19, from_cell=_read_int),
    Column("Home (Home) O1.5 %", "home_home_over15_rate", 18, **_pct()),
    Column("Home (Home) Avg Total", "home_home_avg_total", 20, FORMAT_GOALS, from_cell=_read_float),
    Column("Away (Away) Matches", "away_away_matches", 19, from_cell=_read_int),
    Column("Away (Away) O1.5 %", "away_away_over15_rate", 18, **_pct()),
    Column("Away (Away) Avg Total", "away_away_avg_total", 20, FORMAT_GOALS, from_cell=_read_float),
    Column("League Avg Goals", "league_avg_goals", 16, FORMAT_GOALS, from_cell=_read_float),
    Column("League O1.5 %", "league_over15_rate", 14, **_pct()),
    Column("League Trend", "league_trend", 34, from_cell=_read_str),
    Column("Context Notes", "context_notes", 34, from_cell=_read_str),
    Column("Model Notes", "model_notes", 50, from_cell=_read_str),
    Column("Review Flags", "review_flags", 40, from_cell=_read_str),
    Column("Unreliable Reason", "unreliable_reason", 34, from_cell=_read_str),
    Column("Fixture Source", "fixture_source", 26, from_cell=_read_str),
    Column("Stats Sources", "stats_sources", 26, from_cell=_read_str),
    Column(
        "Collected At (UTC)",
        "collected_at_utc",
        18,
        FORMAT_DATETIME,
        to_cell=_utc_naive,
        from_cell=_read_utc,
    ),
    Column(
        "First Seen (UTC)",
        "first_seen_utc",
        18,
        FORMAT_DATETIME,
        to_cell=_utc_naive,
        from_cell=_read_utc,
    ),
    Column(
        "Last Updated (UTC)",
        "last_updated_utc",
        18,
        FORMAT_DATETIME,
        to_cell=_utc_naive,
        from_cell=_read_utc,
    ),
)

#: ``Ranked Selections`` - exactly the columns section 10.4 tabulates.
RANKED_COLUMNS: tuple[Column, ...] = (
    Column("Rank", "rank", 7),
    Column("Match", "", 30),
    Column("League", "league", 24),
    Column("Kick-off (EAT)", "kickoff_utc", 18, FORMAT_DATETIME, to_cell=_eat_naive),
    Column("O1.5 Odds", "odds_over_15", 10, FORMAT_ODDS),
    Column("Model Prob", "model_probability", 12, FORMAT_PERCENT),
    Column("Implied (Fair)", "implied_fair", 13, FORMAT_PERCENT),
    Column("Value", "value", 10, FORMAT_SIGNED_PERCENT),
    Column("EV", "ev", 10, FORMAT_SIGNED_PERCENT),
    Column("H2H O1.5 %", "h2h_over15_rate", 12, FORMAT_PERCENT),
    Column("Recent O1.5 %", "", 14, FORMAT_PERCENT),
    Column("Confidence", "confidence", 12),
    Column("Category", "category", 20),
    Column("Score", "score", 9, FORMAT_SCORE),
    Column("Status", "status", 22),
    Column("First Seen", "first_seen_utc", 18, FORMAT_DATETIME, to_cell=_utc_naive),
    Column("Last Updated", "last_updated_utc", 18, FORMAT_DATETIME, to_cell=_utc_naive),
)

RUN_LOG_COLUMNS: tuple[tuple[str, str, int, str | None], ...] = (
    ("Run Started (UTC)", "started_at", 18, FORMAT_DATETIME),
    ("Run Finished (UTC)", "finished_at", 18, FORMAT_DATETIME),
    ("Duration (s)", "duration_seconds", 12, "0.00"),
    ("Window (EAT)", "window", 18, None),
    ("Min Probability", "min_probability", 15, "0.0"),
    ("Min Odds", "min_odds", 10, FORMAT_ODDS),
    ("Fixtures Found", "fixtures_found", 14, None),
    ("Rows Added", "rows_added", 12, None),
    ("Rows Updated", "rows_updated", 13, None),
    ("Warnings", "warnings", 10, None),
    ("Errors", "errors", 9, None),
    ("Warning Detail", "warning_detail", 60, None),
    ("Error Detail", "error_detail", 60, None),
    ("Effective Configuration", "config_summary", 80, None),
)


def recent_over15_rate(record: MatchRecord) -> float | None:
    """Mean of the two sides' recent Over 1.5 rates, where available."""
    rates = [
        rate
        for rate in (record.home_form_over15_rate, record.away_form_over15_rate)
        if rate is not None
    ]
    return sum(rates) / len(rates) if rates else None


def _cell_value(record: MatchRecord, column: Column) -> Any:
    """Value for one cell, including the derived columns."""
    if column.header == "Match":
        return record.label
    if column.header == "Recent O1.5 %":
        return recent_over15_rate(record)
    if not column.attribute:
        return None
    value = getattr(record, column.attribute, None)
    return column.to_cell(value) if column.to_cell else value


# ==========================================================================
# Merge (spec 10.2)
# ==========================================================================
@dataclass
class MergeOutcome:
    """The merged record set plus the counts the Run Log reports."""

    records: list[MatchRecord]
    added: int = 0
    updated: int = 0
    retained_absent: int = 0


#: Fields describing the fixture and its markets - always refreshed.
_MARKET_FIELDS = (
    "league",
    "home_team",
    "away_team",
    "odds_over_15",
    "odds_under_15",
    "odds_over_25",
    "odds_under_25",
    "odds_home",
    "odds_draw",
    "odds_away",
    "odds_btts_yes",
    "odds_btts_no",
    "collected_at_utc",
    "fixture_source",
)


def merge_records(
    existing: Sequence[MatchRecord],
    fresh: Sequence[MatchRecord],
    *,
    timestamp: datetime | None = None,
) -> MergeOutcome:
    """Upsert *fresh* onto *existing* on ``match_key`` (spec 10.2).

    ============================================  ===========================
    Situation                                     Behaviour
    ============================================  ===========================
    Key already in the workbook                   Overwrite with new values
    Key is new                                    Append it
    Key present, collected but not analysed       Refresh markets, keep the
                                                  earlier analysis
    Key in the workbook, absent from this run     Retain, mark, leave
                                                  ``Last Updated`` untouched
    ============================================  ===========================

    Rows are never deleted.
    """
    moment = timestamp or now_utc()
    by_key = {record.match_key: record for record in existing}
    outcome = MergeOutcome(records=[])
    seen: set[str] = set()

    for record in fresh:
        seen.add(record.match_key)
        previous = by_key.get(record.match_key)
        if previous is None:
            record.first_seen_utc = record.first_seen_utc or moment
            record.last_updated_utc = moment
            by_key[record.match_key] = record
            outcome.added += 1
            continue

        record.first_seen_utc = previous.first_seen_utc or moment
        record.last_updated_utc = moment
        if record.model_probability is None and previous.model_probability is not None:
            # Collected but not analysed this run - keep the earlier analysis
            # and refresh only what was actually re-observed.
            for field_name in _MARKET_FIELDS:
                setattr(previous, field_name, getattr(record, field_name))
            previous.last_updated_utc = moment
            by_key[record.match_key] = previous
        else:
            by_key[record.match_key] = record
        outcome.updated += 1

    for key, record in by_key.items():
        if key not in seen:
            record.status = STATUS_NOT_IN_LATEST_RUN
            outcome.retained_absent += 1  # Last Updated deliberately untouched

    outcome.records = list(by_key.values())
    return outcome


# ==========================================================================
# The store
# ==========================================================================
class ExcelStore:
    """Reads and writes the one workbook that belongs to a match date."""

    def __init__(self, match_day: date, config: Config) -> None:
        self.match_day = match_day
        self.config = config
        self.directory = Path(config.output_dir)
        self.path = self.directory / f"{match_day.isoformat()}.xlsx"

    # -- reading -------------------------------------------------------
    def exists(self) -> bool:
        return self.path.is_file()

    def load(self) -> list[MatchRecord]:
        """Read ``All Matches`` back into records. A missing file yields []."""
        if not self.exists():
            return []
        try:
            workbook = load_workbook(self.path, data_only=True)
        except Exception as exc:  # noqa: BLE001 - a damaged file must not abort
            log.error(
                "could not read the existing workbook",
                extra={"path": str(self.path), "error": str(exc)},
            )
            return []
        try:
            if SHEET_ALL not in workbook.sheetnames:
                return []
            sheet = workbook[SHEET_ALL]
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = [str(h).strip() if h is not None else "" for h in next(rows)]
            except StopIteration:
                return []

            by_header = {column.header: column for column in ALL_MATCHES_COLUMNS}
            records: list[MatchRecord] = []
            for raw in rows:
                # Not strict: a workbook written by an older version may have
                # fewer columns than the current schema, and that must still load.
                values = dict(zip(headers, raw))  # noqa: B905
                key = _read_str(values.get("Match Key")).strip()
                if not key:
                    continue
                kickoff = _read_utc(values.get("Kick-off (UTC)"))
                if kickoff is None:
                    naive_eat = values.get("Kick-off (EAT)")
                    if isinstance(naive_eat, datetime):
                        from models import EAT

                        kickoff = to_utc(naive_eat.replace(tzinfo=EAT))
                if kickoff is None:
                    log.warning("skipping stored row with no readable kick-off", extra={"key": key})
                    continue

                record = MatchRecord(
                    match_key=key,
                    league=_read_str(values.get("League")),
                    home_team=_read_str(values.get("Home Team")),
                    away_team=_read_str(values.get("Away Team")),
                    kickoff_utc=kickoff,
                )
                for header, column in by_header.items():
                    if header in {"Match Key", "Match", "Kick-off (EAT)", "Kick-off (UTC)"}:
                        continue
                    if not column.attribute or column.from_cell is None:
                        continue
                    setattr(record, column.attribute, column.from_cell(values.get(header)))
                records.append(record)

            log.info(
                "loaded existing workbook", extra={"path": str(self.path), "rows": len(records)}
            )
            return records
        finally:
            workbook.close()

    def first_run_timestamp(self) -> datetime | None:
        """When this workbook was first written, from the Run Log."""
        if not self.exists():
            return None
        try:
            workbook = load_workbook(self.path, data_only=True)
        except Exception:  # noqa: BLE001
            return None
        try:
            if SHEET_RUN_LOG not in workbook.sheetnames:
                return None
            rows = workbook[SHEET_RUN_LOG].iter_rows(min_row=2, max_row=2, values_only=True)
            for row in rows:
                return _read_utc(row[0]) if row else None
        finally:
            workbook.close()
        return None

    def existing_run_log(self) -> list[tuple]:
        """Previous Run Log rows, so history accumulates across runs."""
        if not self.exists():
            return []
        try:
            workbook = load_workbook(self.path, data_only=True)
        except Exception:  # noqa: BLE001
            return []
        try:
            if SHEET_RUN_LOG not in workbook.sheetnames:
                return []
            rows = list(workbook[SHEET_RUN_LOG].iter_rows(min_row=2, values_only=True))
            return [row for row in rows if any(cell is not None for cell in row)]
        finally:
            workbook.close()

    # -- writing -------------------------------------------------------
    def save(
        self,
        records: Sequence[MatchRecord],
        *,
        window: AnalysisWindow,
        summary_rows: Sequence[tuple[str, Any]],
        run_log: RunLogEntry,
        sources: Sequence[SourceRecord],
    ) -> Path:
        """Write every owned sheet and move the result onto the target file."""
        self.directory.mkdir(parents=True, exist_ok=True)
        self._raise_if_locked()

        workbook = load_workbook(self.path) if self.exists() else Workbook()
        try:
            if not self.exists() and "Sheet" in workbook.sheetnames:
                del workbook["Sheet"]

            previous_run_log = self.existing_run_log()
            for name in OWNED_SHEETS:
                if name in workbook.sheetnames:
                    del workbook[name]

            self._write_summary(workbook.create_sheet(SHEET_SUMMARY), summary_rows)
            self._write_ranked(workbook.create_sheet(SHEET_RANKED), records)
            self._write_all_matches(workbook.create_sheet(SHEET_ALL), records)
            self._write_run_log(workbook.create_sheet(SHEET_RUN_LOG), previous_run_log, run_log)
            self._write_sources(workbook.create_sheet(SHEET_SOURCES), sources)

            # Keep the owned sheets in the documented order, ahead of any
            # sheets the user added themselves.
            order = [workbook[name] for name in OWNED_SHEETS if name in workbook.sheetnames]
            others = [s for s in workbook.worksheets if s not in order]
            workbook._sheets = order + others  # noqa: SLF001 - openpyxl has no public API

            self._atomic_save(workbook)
        finally:
            workbook.close()

        log.info("workbook written", extra={"path": str(self.path), "rows": len(records)})
        return self.path

    # -- write safety (spec 10.3) --------------------------------------
    def _lock_file(self) -> Path:
        return self.path.with_name(f"~${self.path.name}")

    def _raise_if_locked(self) -> None:
        """Fail early and clearly when the workbook is open in Excel."""
        if not self.exists():
            return
        if self._lock_file().is_file():
            log.warning("an Excel lock file is present", extra={"lock": str(self._lock_file())})
        try:
            with self.path.open("r+b"):
                return
        except PermissionError as exc:
            raise WorkbookLockedError(self._locked_message()) from exc
        except OSError:
            return

    def _locked_message(self) -> str:
        return (
            f"cannot write {self.path}: the file is open in another program.\n"
            f"  Close {self.path} in Excel and re-run. Cached results make the\n"
            f"  re-run inexpensive.\n"
            f"  The tool will not write a differently named file: there must be\n"
            f"  exactly one workbook per date."
        )

    def _atomic_save(self, workbook: Workbook) -> None:
        """Build beside the target, then ``os.replace`` onto it (spec 10.3)."""
        temporary = self.path.with_name(f".{self.path.stem}.{os.getpid()}.tmp.xlsx")
        workbook.save(temporary)

        delay = 0.5
        last_error: Exception | None = None
        for attempt in range(1, max(1, self.config.excel_lock_retries) + 1):
            try:
                os.replace(temporary, self.path)
                return
            except PermissionError as exc:
                last_error = exc
                if attempt >= self.config.excel_lock_retries:
                    break
                log.warning(
                    "target workbook is locked; retrying",
                    extra={"attempt": attempt, "delay_seconds": delay, "path": str(self.path)},
                )
                time_module.sleep(delay)
                delay *= self.config.excel_lock_backoff

        temporary.unlink(missing_ok=True)
        raise WorkbookLockedError(self._locked_message()) from last_error

    # -- sheet builders -------------------------------------------------
    @staticmethod
    def _style_header(sheet: Worksheet, columns: Sequence[str], widths: Sequence[int]) -> None:
        for index, (header, width) in enumerate(zip(columns, widths, strict=True), start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        if columns:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(sheet.max_row, 1)}"

    def _write_table(
        self,
        sheet: Worksheet,
        columns: Sequence[Column],
        records: Sequence[MatchRecord],
        colour_scale_headers: Iterable[str] = (),
    ) -> None:
        for row_index, record in enumerate(records, start=2):
            for column_index, column in enumerate(columns, start=1):
                cell = sheet.cell(
                    row=row_index, column=column_index, value=_cell_value(record, column)
                )
                if column.number_format:
                    cell.number_format = column.number_format

        self._style_header(sheet, [c.header for c in columns], [c.width for c in columns])

        if records:
            for header in colour_scale_headers:
                index = next(
                    (i for i, c in enumerate(columns, start=1) if c.header == header), None
                )
                if index is None:
                    continue
                letter = get_column_letter(index)
                sheet.conditional_formatting.add(
                    f"{letter}2:{letter}{len(records) + 1}",
                    ColorScaleRule(
                        start_type="min",
                        start_color="F8696B",
                        mid_type="percentile",
                        mid_value=50,
                        mid_color="FFEB84",
                        end_type="max",
                        end_color="63BE7B",
                    ),
                )

    def _write_ranked(self, sheet: Worksheet, records: Sequence[MatchRecord]) -> None:
        """Rebuilt from the full retained set every run (spec 2.6)."""
        selections = ranked_selections(list(records))
        self._write_table(sheet, RANKED_COLUMNS, selections, ("Model Prob", "Value"))
        if not selections:
            sheet.cell(
                row=2,
                column=1,
                value="No matches qualified under the thresholds used for this run.",
            )
            sheet.cell(
                row=3,
                column=1,
                value=(
                    "See the Summary sheet for the thresholds, and All Matches "
                    "for the reason each fixture was excluded."
                ),
            )

    def _write_all_matches(self, sheet: Worksheet, records: Sequence[MatchRecord]) -> None:
        ordered = sorted(
            records,
            key=lambda r: (
                r.rank if r.rank is not None else 10**6,
                -(r.model_probability or 0.0),
                r.kickoff_utc,
            ),
        )
        self._write_table(sheet, ALL_MATCHES_COLUMNS, ordered, ("Model Prob", "Value"))

    def _write_summary(self, sheet: Worksheet, rows: Sequence[tuple[str, Any]]) -> None:
        sheet.cell(row=1, column=1, value="Over 1.5 Goals - Match Analysis").font = TITLE_FONT
        current = 3
        for label, value in rows:
            if label == "":
                current += 1
                continue
            if value is None:
                cell = sheet.cell(row=current, column=1, value=label)
                cell.font = SECTION_FONT
            else:
                sheet.cell(row=current, column=1, value=label).font = SECTION_FONT
                sheet.cell(row=current, column=2, value=value)
            current += 1

        current += 1
        sheet.cell(row=current, column=1, value="Disclaimer").font = SECTION_FONT
        note = sheet.cell(row=current + 1, column=1, value=DISCLAIMER)
        note.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=current + 1, start_column=1, end_row=current + 3, end_column=4)

        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 46
        sheet.column_dimensions["C"].width = 18
        sheet.column_dimensions["D"].width = 18

    def _write_run_log(
        self,
        sheet: Worksheet,
        previous: Sequence[tuple],
        entry: RunLogEntry,
    ) -> None:
        headers = [h for h, _, _, _ in RUN_LOG_COLUMNS]
        widths = [w for _, _, w, _ in RUN_LOG_COLUMNS]

        row_index = 2
        for row in previous:
            for column_index, value in enumerate(row[: len(headers)], start=1):
                sheet.cell(row=row_index, column=column_index, value=value)
            row_index += 1

        for column_index, (_, attribute, _, number_format) in enumerate(RUN_LOG_COLUMNS, start=1):
            value = getattr(entry, attribute, None)
            if isinstance(value, datetime):
                value = _utc_naive(value)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if number_format:
                cell.number_format = number_format

        self._style_header(sheet, headers, widths)

    def _write_sources(self, sheet: Worksheet, sources: Sequence[SourceRecord]) -> None:
        headers = ["Statistic", "Source", "Endpoint / URL", "Retrieved (UTC)", "Notes"]
        widths = [28, 26, 60, 20, 60]
        for row_index, record in enumerate(sources, start=2):
            sheet.cell(row=row_index, column=1, value=record.statistic)
            sheet.cell(row=row_index, column=2, value=record.source)
            sheet.cell(row=row_index, column=3, value=record.endpoint)
            cell = sheet.cell(row=row_index, column=4, value=_utc_naive(record.retrieved_at))
            cell.number_format = FORMAT_DATETIME
            sheet.cell(row=row_index, column=5, value=record.notes)
        self._style_header(sheet, headers, widths)


# ==========================================================================
# Summary assembly
# ==========================================================================
def build_summary_rows(
    *,
    match_day: date,
    window: AnalysisWindow,
    config: Config,
    first_run: datetime | None,
    latest_run: datetime,
    found: int,
    in_window: int,
    analysed: int,
    qualified: int,
    by_category: dict[str, int],
    fixture_source: str,
) -> list[tuple[str, Any]]:
    """The ``Summary`` sheet contents, including the source of every setting."""
    rows: list[tuple[str, Any]] = [
        ("Analysis date", match_day.isoformat()),
        ("Analysis window (EAT)", f"{window.label} ({window.kind})"),
        ("First run (UTC)", _utc_naive(first_run or latest_run)),
        ("Latest run (UTC)", _utc_naive(latest_run)),
        ("Fixture source", fixture_source),
        ("", None),
        ("Matches found", found),
        ("Matches in window", in_window),
        ("Matches analysed", analysed),
        ("Matches qualified", qualified),
        ("", None),
    ]
    for category in ("TOP PICK", "GOOD PICK", "HIGH-ODDS VALUE PICK"):
        rows.append((f"  {category}", by_category.get(category, 0)))
    rows.append(("", None))
    rows.append(("Effective configuration", None))
    for name, value, source in config.effective_summary():
        rows.append((f"  {name}", f"{value}   [{source}]"))
    return rows
