"""The Excel upsert path and window filtering (spec 10; testing bullets 8, 9).

Between them these cover:

* fixtures outside the window are excluded from the recommendations but kept
  in ``All Matches``
* creating a new workbook, and updating an existing one in place
* retaining rows absent from a later run
* rebuilding ``Ranked Selections`` when the thresholds change between runs
* never producing a second file for a date
"""

from __future__ import annotations

from datetime import date, datetime, time, timedelta

import pytest
from conftest import MATCH_DAY, eat, make_record
from openpyxl import Workbook, load_workbook

from excel_writer import (
    SHEET_ALL,
    SHEET_RANKED,
    SHEET_RUN_LOG,
    SHEET_SOURCES,
    SHEET_SUMMARY,
    ExcelStore,
    WorkbookLockedError,
    build_summary_rows,
    merge_records,
)
from models import (
    STATUS_ANALYSED,
    STATUS_NOT_IN_LATEST_RUN,
    STATUS_OUTSIDE_WINDOW,
    UTC,
    AnalysisWindow,
    RunLogEntry,
    SourceRecord,
)
from ranking import apply_window, rank_and_categorise, ranked_selections


@pytest.fixture
def store(config, tmp_path) -> ExcelStore:
    config.output_dir = str(tmp_path)
    return ExcelStore(MATCH_DAY, config)


def save(store, records, *, window=None, config=None, sources=()):
    """Persist *records* the way ``main`` does."""
    window = window or AnalysisWindow.full_day(MATCH_DAY)
    now = datetime(2026, 9, 7, 12, 0, tzinfo=UTC)
    entry = RunLogEntry(
        started_at=now,
        finished_at=now,
        duration_seconds=0.1,
        window=window.label,
        min_probability=(config or store.config).min_probability,
        min_odds=(config or store.config).min_odds,
        fixtures_found=len(records),
        rows_added=0,
        rows_updated=0,
        warnings=0,
        errors=0,
    )
    rows = build_summary_rows(
        match_day=MATCH_DAY,
        window=window,
        config=config or store.config,
        first_run=store.first_run_timestamp(),
        latest_run=now,
        found=len(records),
        in_window=len(records),
        analysed=len(records),
        qualified=sum(1 for r in records if r.qualified),
        by_category={},
        fixture_source="test",
    )
    return store.save(
        records, window=window, summary_rows=rows, run_log=entry, sources=list(sources)
    )


# ==========================================================================
# Window filtering (testing bullet 8)
# ==========================================================================
def test_a_fixture_outside_the_window_is_marked_not_dropped(config):
    inside = make_record("A", "B", hour=15)
    outside = make_record("C", "D", hour=23)
    window = AnalysisWindow.from_times(MATCH_DAY, time(14, 0), time(18, 0))

    apply_window([inside, outside], window)
    assert inside.status == STATUS_ANALYSED
    assert outside.status == STATUS_OUTSIDE_WINDOW


def test_an_out_of_window_fixture_is_excluded_from_recommendations_but_retained(config):
    inside = make_record("A", "B", hour=15)
    outside = make_record("C", "D", hour=23)
    window = AnalysisWindow.from_times(MATCH_DAY, time(14, 0), time(18, 0))

    records = [inside, outside]
    summary = rank_and_categorise(records, config, window)

    assert summary.qualified == 1
    assert [r.match_key for r in ranked_selections(records)] == [inside.match_key]
    assert outside in records  # still in the full set
    assert outside.exclusion_reason


def test_widening_the_window_brings_a_fixture_back(config):
    """Re-marking matters as much as marking."""
    record = make_record("C", "D", hour=23)
    narrow = AnalysisWindow.from_times(MATCH_DAY, time(14, 0), time(18, 0))
    rank_and_categorise([record], config, narrow)
    assert record.status == STATUS_OUTSIDE_WINDOW

    rank_and_categorise([record], config, AnalysisWindow.full_day(MATCH_DAY))
    assert record.status == STATUS_ANALYSED
    assert record.qualified is True


def test_window_bounds_are_inclusive():
    window = AnalysisWindow.from_times(MATCH_DAY, time(14, 0), time(18, 0))
    assert window.contains(eat(14, 0))
    assert window.contains(eat(18, 0))
    assert not window.contains(eat(13, 59))
    assert not window.contains(eat(18, 1))


def test_the_full_day_window_covers_midnight_to_end_of_day():
    window = AnalysisWindow.full_day(MATCH_DAY)
    assert window.contains(eat(0, 0))
    assert window.contains(eat(23, 59))
    assert not window.contains(eat(0, 0, MATCH_DAY + timedelta(days=1)))


def test_out_of_window_rows_reach_the_all_matches_sheet(store, config):
    inside = make_record("A", "B", hour=15)
    outside = make_record("C", "D", hour=23)
    window = AnalysisWindow.from_times(MATCH_DAY, time(14, 0), time(18, 0))
    rank_and_categorise([inside, outside], config, window)

    path = save(store, [inside, outside], window=window)
    workbook = load_workbook(path)
    try:
        assert workbook[SHEET_ALL].max_row == 3  # header plus both fixtures
        assert workbook[SHEET_RANKED].max_row == 2  # header plus the one selection
        statuses = {row[7] for row in workbook[SHEET_ALL].iter_rows(min_row=2, values_only=True)}
        assert STATUS_OUTSIDE_WINDOW in statuses
    finally:
        workbook.close()


# ==========================================================================
# Merge semantics (spec 10.2)
# ==========================================================================
def test_a_new_key_is_appended():
    outcome = merge_records([], [make_record("A", "B")])
    assert outcome.added == 1 and outcome.updated == 0
    assert len(outcome.records) == 1


def test_an_existing_key_is_overwritten_not_duplicated():
    first = make_record("A", "B", probability=0.75)
    second = make_record("A", "B", probability=0.88)
    assert first.match_key == second.match_key

    outcome = merge_records([first], [second])
    assert outcome.added == 0 and outcome.updated == 1
    assert len(outcome.records) == 1
    assert outcome.records[0].model_probability == pytest.approx(0.88)


def test_first_seen_survives_an_update_and_last_updated_moves():
    earlier = datetime(2026, 9, 7, 8, 0, tzinfo=UTC)
    later = datetime(2026, 9, 7, 20, 0, tzinfo=UTC)

    first = make_record("A", "B")
    merge_records([], [first], timestamp=earlier)
    assert first.first_seen_utc == earlier

    second = make_record("A", "B", probability=0.9)
    outcome = merge_records([first], [second], timestamp=later)
    assert outcome.records[0].first_seen_utc == earlier
    assert outcome.records[0].last_updated_utc == later


def test_a_row_absent_from_this_run_is_retained_and_marked():
    """Spec 10.2 - a postponed fixture stays visible with its history."""
    existing = make_record("A", "B")
    existing.last_updated_utc = datetime(2026, 9, 7, 8, 0, tzinfo=UTC)

    outcome = merge_records([existing], [make_record("C", "D")])
    assert len(outcome.records) == 2
    assert outcome.retained_absent == 1

    retained = next(r for r in outcome.records if r.home_team == "A")
    assert retained.status == STATUS_NOT_IN_LATEST_RUN
    assert retained.last_updated_utc == datetime(2026, 9, 7, 8, 0, tzinfo=UTC)


def test_a_collected_but_unanalysed_fixture_keeps_its_earlier_analysis():
    analysed = make_record("A", "B", probability=0.85, odds=1.40)
    fresh = make_record("A", "B", probability=None, odds=1.55)
    fresh.model_probability = None

    outcome = merge_records([analysed], [fresh])
    kept = outcome.records[0]
    assert kept.model_probability == pytest.approx(0.85)  # analysis preserved
    assert kept.odds_over_15 == pytest.approx(1.55)  # market refreshed


def test_rows_are_never_deleted():
    existing = [make_record(f"Home{i}", "Away") for i in range(5)]
    outcome = merge_records(existing, [])
    assert len(outcome.records) == 5
    assert all(r.status == STATUS_NOT_IN_LATEST_RUN for r in outcome.records)


# ==========================================================================
# One file per date (spec 10.1, 10.2)
# ==========================================================================
def test_the_filename_is_the_match_date_alone(config, tmp_path):
    config.output_dir = str(tmp_path)
    assert ExcelStore(date(2026, 9, 7), config).path.name == "2026-09-07.xlsx"
    assert ExcelStore(date(2026, 12, 31), config).path.name == "2026-12-31.xlsx"


def test_the_filename_ignores_the_window_and_the_thresholds(config, tmp_path):
    """Spec 10.1 - a narrower window still writes to the same file."""
    config.output_dir = str(tmp_path)
    first = ExcelStore(MATCH_DAY, config).path

    config.min_probability = 90.0
    config.window_start, config.window_end = "19:00", "21:00"
    assert ExcelStore(MATCH_DAY, config).path == first


def test_a_missing_output_directory_is_created(config, tmp_path):
    config.output_dir = str(tmp_path / "nested" / "results")
    store = ExcelStore(MATCH_DAY, config)
    save(store, [make_record()])
    assert store.path.is_file()


def test_the_first_run_creates_the_workbook_with_every_sheet(store):
    path = save(store, [make_record()])
    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames[:5] == [
            SHEET_SUMMARY,
            SHEET_RANKED,
            SHEET_ALL,
            SHEET_RUN_LOG,
            SHEET_SOURCES,
        ]
    finally:
        workbook.close()


def test_re_running_updates_the_one_file_and_creates_no_second(store, config):
    """Acceptance criterion 10, the headline rule of section 10.2."""
    first = [make_record("A", "B")]
    rank_and_categorise(first, config)
    save(store, first)

    fresh = [make_record("A", "B", probability=0.90), make_record("C", "D")]
    merged = merge_records(store.load(), fresh)
    rank_and_categorise(merged.records, config)
    save(store, merged.records)

    files = sorted(p.name for p in store.directory.glob("*.xlsx"))
    assert files == ["2026-09-08.xlsx"]
    assert merged.added == 1 and merged.updated == 1


def test_no_suffixed_or_timestamped_variant_is_ever_written(store, config):
    for probability in (0.75, 0.80, 0.85, 0.90):
        records = [make_record("A", "B", probability=probability)]
        merged = merge_records(store.load(), records)
        rank_and_categorise(merged.records, config)
        save(store, merged.records)

    assert [p.name for p in store.directory.glob("*.xlsx")] == ["2026-09-08.xlsx"]
    assert list(store.directory.glob("*(1)*")) == []
    assert list(store.directory.glob("*_v2*")) == []
    assert list(store.directory.glob("*.tmp.xlsx")) == []


def test_a_round_trip_preserves_the_figures(store, config):
    original = make_record("A", "B", probability=0.842, odds=1.44, value=0.061)
    original.lambda_total = 3.30
    original.h2h_matches = 8
    original.h2h_over15_rate = 0.80
    original.implied_fair = 0.781
    rank_and_categorise([original], config)
    save(store, [original])

    loaded = store.load()
    assert len(loaded) == 1
    restored = loaded[0]
    assert restored.match_key == original.match_key
    assert restored.model_probability == pytest.approx(0.842)
    assert restored.odds_over_15 == pytest.approx(1.44)
    assert restored.value == pytest.approx(0.061)
    assert restored.lambda_total == pytest.approx(3.30)
    assert restored.h2h_matches == 8
    assert restored.kickoff_utc == original.kickoff_utc
    assert restored.confidence == original.confidence


def test_loading_a_workbook_that_does_not_exist_yields_nothing(store):
    assert store.exists() is False
    assert store.load() == []


# ==========================================================================
# Rebuilding under changed thresholds (spec 2.6)
# ==========================================================================
def test_ranked_selections_is_rebuilt_when_the_thresholds_tighten(store, config):
    """A fixture that no longer qualifies moves back to All Matches."""
    records = [make_record("A", "B", probability=0.75), make_record("C", "D", probability=0.88)]
    rank_and_categorise(records, config)
    save(store, records)

    workbook = load_workbook(store.path)
    assert workbook[SHEET_RANKED].max_row == 3  # header plus both
    workbook.close()

    config.min_probability = 85.0
    reloaded = store.load()
    rank_and_categorise(reloaded, config)
    save(store, reloaded, config=config)

    workbook = load_workbook(store.path)
    try:
        assert workbook[SHEET_RANKED].max_row == 2  # header plus the survivor
        assert workbook[SHEET_ALL].max_row == 3  # both fixtures still recorded
        reasons = [row[8] for row in workbook[SHEET_ALL].iter_rows(min_row=2, values_only=True)]
        assert any(reason and "Below minimum probability" in reason for reason in reasons)
    finally:
        workbook.close()


def test_loosening_the_threshold_brings_a_fixture_back(store, config):
    records = [make_record("A", "B", probability=0.75)]
    config.min_probability = 85.0
    rank_and_categorise(records, config)
    save(store, records, config=config)
    assert records[0].qualified is False

    config.min_probability = 70.0
    reloaded = store.load()
    rank_and_categorise(reloaded, config)
    assert reloaded[0].qualified is True


def test_an_empty_ranked_sheet_says_so_plainly(store, config):
    records = [make_record("A", "B", probability=0.20, value=-0.4)]
    rank_and_categorise(records, config)
    save(store, records)

    workbook = load_workbook(store.path)
    try:
        assert "No matches qualified" in str(workbook[SHEET_RANKED].cell(row=2, column=1).value)
    finally:
        workbook.close()


# ==========================================================================
# The run log, sources and formatting (spec 10.4, 10.5)
# ==========================================================================
def test_the_run_log_accumulates_a_row_per_execution(store, config):
    for _ in range(3):
        records = store.load() or [make_record()]
        save(store, records)

    workbook = load_workbook(store.path)
    try:
        assert workbook[SHEET_RUN_LOG].max_row == 4  # header plus three runs
    finally:
        workbook.close()


def test_the_summary_records_the_source_of_every_setting(store, config):
    from config import SOURCE_FLAG

    config.set("min_probability", 82.0, SOURCE_FLAG)
    save(store, [make_record()], config=config)

    workbook = load_workbook(store.path)
    try:
        text = "\n".join(
            str(cell.value)
            for row in workbook[SHEET_SUMMARY].iter_rows(values_only=False)
            for cell in row
            if cell.value is not None
        )
        assert "Minimum probability" in text
        assert "82.0%" in text and "[flag]" in text
        assert "statistical estimates" in text  # the disclaimer
    finally:
        workbook.close()


def test_data_sources_carries_per_statistic_provenance(store):
    sources = [
        SourceRecord(
            "A vs B: head-to-head", "test feed", "https://example.test", None, "4 meetings"
        ),
        SourceRecord(
            "A vs B: league baseline", "test feed", "https://example.test", None, "unavailable"
        ),
    ]
    save(store, [make_record()], sources=sources)

    workbook = load_workbook(store.path)
    try:
        rows = list(workbook[SHEET_SOURCES].iter_rows(min_row=2, values_only=True))
        assert len(rows) == 2
        assert rows[0][0] == "A vs B: head-to-head"
        assert rows[1][4] == "unavailable"
    finally:
        workbook.close()


def test_headers_are_frozen_and_filtered(store):
    save(store, [make_record()])
    workbook = load_workbook(store.path)
    try:
        for name in (SHEET_RANKED, SHEET_ALL, SHEET_RUN_LOG, SHEET_SOURCES):
            assert workbook[name].freeze_panes == "A2"
            assert workbook[name].auto_filter.ref
    finally:
        workbook.close()


def test_percentages_and_odds_carry_the_right_number_formats(store, config):
    records = [make_record()]
    rank_and_categorise(records, config)
    save(store, records)

    workbook = load_workbook(store.path)
    try:
        sheet = workbook[SHEET_RANKED]
        headers = [cell.value for cell in sheet[1]]
        probability = sheet.cell(row=2, column=headers.index("Model Prob") + 1)
        odds = sheet.cell(row=2, column=headers.index("O1.5 Odds") + 1)
        assert "%" in probability.number_format
        assert odds.number_format == "0.00"
    finally:
        workbook.close()


def test_sheets_the_tool_does_not_own_are_preserved(store):
    save(store, [make_record()])

    workbook = load_workbook(store.path)
    workbook.create_sheet("My Notes")["A1"] = "keep me"
    workbook.save(store.path)
    workbook.close()

    save(store, [make_record()])

    workbook = load_workbook(store.path)
    try:
        assert "My Notes" in workbook.sheetnames
        assert workbook["My Notes"]["A1"].value == "keep me"
    finally:
        workbook.close()


def test_a_damaged_workbook_is_reported_rather_than_crashing_the_run(store):
    store.directory.mkdir(parents=True, exist_ok=True)
    store.path.write_text("this is not a workbook", encoding="utf-8")
    assert store.load() == []


def test_a_workbook_without_an_all_matches_sheet_loads_as_empty(store):
    store.directory.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.save(store.path)
    workbook.close()
    assert store.load() == []


def test_a_locked_workbook_fails_with_an_actionable_message(store, config, monkeypatch):
    """Spec 10.3 - never fall back to a differently named file."""
    import os

    save(store, [make_record()])
    config.excel_lock_retries = 1

    def refuse(*_args, **_kwargs):
        raise PermissionError("locked by Excel")

    monkeypatch.setattr(os, "replace", refuse)
    with pytest.raises(WorkbookLockedError) as excinfo:
        save(store, [make_record()])

    message = str(excinfo.value)
    assert "open in another program" in message
    assert "exactly one workbook per date" in message
    assert [p.name for p in store.directory.glob("*.xlsx")] == ["2026-09-08.xlsx"]


def test_a_failed_write_leaves_no_temporary_file_behind(store, config, monkeypatch):
    import os

    save(store, [make_record()])
    config.excel_lock_retries = 1
    monkeypatch.setattr(os, "replace", lambda *a, **k: (_ for _ in ()).throw(PermissionError()))

    with pytest.raises(WorkbookLockedError):
        save(store, [make_record()])
    assert list(store.directory.glob("*.tmp.xlsx")) == []
