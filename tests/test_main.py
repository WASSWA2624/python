"""End-to-end runs through ``main`` (spec 11, 12, 13, 14).

Every run here is unattended and offline: fixtures come from a CSV, statistics
from a local results file, and nothing touches the network.
"""

from __future__ import annotations

from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

import main
from main import EXIT_FATAL, EXIT_OK, EXIT_RECOVERABLE, build_parser
from reporting import RunOutcome, build_rationale, build_report

MATCH_DAY = date(2026, 9, 8)

FIXTURES = """League,Home Team,Away Team,Date,Time,Over 1.5,Under 1.5,Over 2.5,1,X,2
Test League,Alpha,Beta,2026-09-08,15:00,1.40,2.90,2.10,2.5,3.4,2.7
Test League,Gamma,Delta,2026-09-08,20:30,1.55,2.45,2.30,2.8,3.3,2.45
Test League,Epsilon,Zeta,2026-09-08,23:30,2.10,1.75,3.20,2.6,3.1,2.8
"""

TEAMS = ["Alpha", "Beta", "Gamma", "Delta", "Epsilon", "Zeta"]


def _results_csv() -> str:
    """Completed results for the six test teams: a goal-rich league."""
    rows = ["date,league,home,away,home_goals,away_goals"]
    day = MATCH_DAY - timedelta(days=7)
    for round_number in range(12):
        for index in range(0, len(TEAMS), 2):
            home = TEAMS[index]
            away = TEAMS[(index + 1 + round_number) % len(TEAMS)]
            if home == away:
                continue
            played = day - timedelta(days=7 * round_number)
            goals_home = 2 if round_number % 2 == 0 else 1
            goals_away = 1 if round_number % 3 else 2
            rows.append(f"{played.isoformat()},Test League,{home},{away},{goals_home},{goals_away}")
    return "\n".join(rows) + "\n"


@pytest.fixture
def workspace(tmp_path):
    (tmp_path / "fixtures.csv").write_text(FIXTURES, encoding="utf-8")
    (tmp_path / "results.csv").write_text(_results_csv(), encoding="utf-8")
    return tmp_path


def run(workspace, *extra, day: str = "2026-09-08") -> int:
    return main.run(
        [
            "--no-prompt",
            "--date",
            day,
            "--fixtures",
            str(workspace / "fixtures.csv"),
            "--results",
            str(workspace / "results.csv"),
            "--stats-providers",
            "none",
            "--output-dir",
            str(workspace / "results"),
            "--no-cache",
            *extra,
        ]
    )


def workbook_at(workspace, day: str = "2026-09-08"):
    return load_workbook(workspace / "results" / f"{day}.xlsx")


# ==========================================================================
# Acceptance criteria 3, 5, 6, 9
# ==========================================================================
def test_an_unattended_run_completes_and_writes_the_workbook(workspace, capsys):
    assert run(workspace) == EXIT_OK
    assert (workspace / "results" / "2026-09-08.xlsx").is_file()

    output = capsys.readouterr().out
    assert "MATCHES FOUND: 3" in output
    assert "IN WINDOW: 3" in output
    assert "statistical estimates" in output  # the disclaimer


def test_the_effective_configuration_is_logged_for_an_unattended_run(workspace, capsys):
    """Acceptance criterion 3 - an unattended run is self-documenting."""
    run(workspace, "--min-probability", "72", "--min-odds", "1.30")
    output = capsys.readouterr().out
    assert "min probability 72.0%" in output
    assert "min odds 1.30" in output
    assert "window 00:00 - 23:59 EAT" in output


def test_every_fixture_is_either_estimated_or_excluded_with_a_reason(workspace):
    """Acceptance criterion 6."""
    run(workspace)
    sheet = workbook_at(workspace)["All Matches"]
    headers = [cell.value for cell in sheet[1]]
    probability = headers.index("Model Prob")
    reason = headers.index("Exclusion Reason")

    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3
    for row in rows:
        assert row[probability] is not None or row[reason]


def test_the_workbook_is_named_by_the_match_date_alone(workspace):
    """Acceptance criterion 9."""
    run(workspace, "--window", "14:00-22:00", "--min-probability", "90")
    assert [p.name for p in (workspace / "results").glob("*.xlsx")] == ["2026-09-08.xlsx"]


# ==========================================================================
# Acceptance criteria 10 and 11
# ==========================================================================
def test_re_running_updates_the_same_workbook_and_makes_no_second_file(workspace):
    """Acceptance criterion 10, including differing windows and thresholds."""
    assert run(workspace) == EXIT_OK
    assert run(workspace, "--window", "19:00-22:00", "--min-probability", "85") == EXIT_OK
    assert run(workspace, "--min-odds", "1.50") == EXIT_OK

    files = [p.name for p in (workspace / "results").glob("*.xlsx")]
    assert files == ["2026-09-08.xlsx"]

    log = workbook_at(workspace)["Run Log"]
    assert log.max_row == 4  # header plus three runs


def test_narrowing_the_window_retains_rows_and_marks_them(workspace):
    """Acceptance criterion 11 - nothing is deleted when the window narrows."""
    run(workspace)
    run(workspace, "--window", "14:00-16:00")

    sheet = workbook_at(workspace)["All Matches"]
    headers = [cell.value for cell in sheet[1]]
    status = headers.index("Status")
    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    assert len(rows) == 3  # all three still present
    statuses = [row[status] for row in rows]
    assert statuses.count("Outside analysis window") == 2


def test_a_fixture_absent_from_a_later_run_is_retained(workspace):
    run(workspace)

    # The sportsbook drops one fixture on the second run.
    trimmed = "\n".join(FIXTURES.splitlines()[:3]) + "\n"
    (workspace / "fixtures.csv").write_text(trimmed, encoding="utf-8")
    run(workspace)

    sheet = workbook_at(workspace)["All Matches"]
    headers = [cell.value for cell in sheet[1]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3
    assert "Not in latest run" in [row[headers.index("Status")] for row in rows]


def test_the_window_only_narrows_the_recommendations_not_the_collection(workspace, capsys):
    """Acceptance criterion 5 - the full day is collected first."""
    run(workspace, "--window", "14:00-16:00")
    output = capsys.readouterr().out
    assert "MATCHES FOUND: 3" in output
    assert "IN WINDOW: 1" in output


# ==========================================================================
# Acceptance criteria 12, 13, 14
# ==========================================================================
def test_a_fixture_with_no_statistics_is_excluded_not_invented(workspace, capsys):
    """Acceptance criterion 12."""
    (workspace / "fixtures.csv").write_text(
        "League,Home Team,Away Team,Date,Time,Over 1.5,Under 1.5\n"
        "Unknown League,Nowhere United,Nobody City,2026-09-08,18:00,1.40,2.90\n",
        encoding="utf-8",
    )
    assert run(workspace) == EXIT_OK

    sheet = workbook_at(workspace)["All Matches"]
    headers = [cell.value for cell in sheet[1]]
    row = next(sheet.iter_rows(min_row=2, values_only=True))
    assert row[headers.index("Model Prob")] is None
    assert "insufficient historical data" in row[headers.index("Exclusion Reason")].lower()
    assert "NO MATCHES QUALIFIED" in capsys.readouterr().out


def test_one_failing_fixture_does_not_abort_the_run(workspace, monkeypatch, capsys):
    """Acceptance criterion 13."""
    real = main.analyse_fixture
    calls = {"n": 0}

    def sometimes_explode(fixture, provider, config):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RuntimeError("synthetic failure")
        return real(fixture, provider, config)

    monkeypatch.setattr(main, "analyse_fixture", sometimes_explode)

    exit_code = run(workspace)
    assert exit_code == EXIT_RECOVERABLE  # errors logged, run completed
    sheet = workbook_at(workspace)["All Matches"]
    assert sheet.max_row == 4  # header plus all three fixtures
    assert "Errors: 1" in capsys.readouterr().out


def test_no_fixture_source_reaches_the_sportsbook_without_opting_in(workspace, capsys):
    """Acceptance criterion 14 - and the section 0 access policy."""
    exit_code = main.run(
        [
            "--no-prompt",
            "--date",
            "2026-09-08",
            "--output-dir",
            str(workspace / "results"),
            "--stats-providers",
            "none",
            "--no-cache",
        ]
    )
    assert exit_code == EXIT_FATAL
    output = capsys.readouterr().out
    assert "off by default" in output
    assert "--fixtures" in output


# ==========================================================================
# Exit codes and argument handling (spec 12, 13)
# ==========================================================================
def test_a_clean_run_exits_zero(workspace):
    assert run(workspace) == EXIT_OK


def test_a_bad_date_is_a_fatal_argument_error(workspace, capsys):
    assert run(workspace, day="not-a-date") == EXIT_FATAL
    assert "--date must be YYYY-MM-DD" in capsys.readouterr().out


def test_a_missing_fixtures_file_is_fatal(workspace, capsys):
    exit_code = main.run(
        [
            "--no-prompt",
            "--fixtures",
            str(workspace / "absent.csv"),
            "--output-dir",
            str(workspace / "results"),
            "--stats-providers",
            "none",
        ]
    )
    assert exit_code == EXIT_FATAL
    assert "not found" in capsys.readouterr().out


def test_invalid_configuration_is_reported_before_any_work(workspace, capsys, monkeypatch):
    monkeypatch.setenv("OVER15_W_PROBABILITY", "0.9")
    monkeypatch.setenv("OVER15_W_VALUE", "0.4")
    assert run(workspace) == EXIT_FATAL
    assert "sum to 1.0" in capsys.readouterr().out


def test_dry_run_writes_nothing(workspace, capsys):
    assert run(workspace, "--dry-run") == EXIT_OK
    assert not (workspace / "results" / "2026-09-08.xlsx").exists()
    assert "Dry run: no workbook was written." in capsys.readouterr().out


def test_the_window_flags_are_mutually_exclusive():
    with pytest.raises(SystemExit):
        build_parser().parse_args(["--window", "10:00-12:00", "--next-hours", "6"])


def test_every_documented_flag_is_accepted():
    """Every invocation in the specification's command-line section."""
    parser = build_parser()
    for argv in (
        [],
        ["--no-prompt"],
        ["--window", "14:00-22:00"],
        ["--next-hours", "6"],
        ["--min-probability", "75"],
        ["--min-odds", "1.35"],
        ["--date", "2026-09-08"],
        ["--fixtures", "data.csv"],
        ["--output-dir", "results"],
        ["--dry-run"],
    ):
        assert parser.parse_args(argv) is not None


def test_the_dixon_coles_flag_takes_an_optional_rho():
    parser = build_parser()
    assert parser.parse_args(["--dixon-coles"]).dixon_coles == pytest.approx(0.08)
    assert parser.parse_args(["--dixon-coles", "0.05"]).dixon_coles == pytest.approx(0.05)


# ==========================================================================
# The console report (spec 11)
# ==========================================================================
def test_the_report_restates_the_settings_actually_used(workspace, capsys):
    run(workspace, "--window", "14:00-22:00", "--min-probability", "72", "--min-odds", "1.30")
    output = capsys.readouterr().out
    assert "DATE: 8 September 2026" in output
    assert "WINDOW: 14:00 - 22:00 EAT" in output
    assert "THRESHOLDS: min probability 72.0% | min odds 1.30" in output


def test_the_report_names_the_binding_constraint_when_nothing_qualifies(workspace, capsys):
    """Spec 11 - an over-tightened threshold must be visible immediately."""
    run(workspace, "--min-probability", "99")
    output = capsys.readouterr().out
    assert "NO MATCHES QUALIFIED." in output
    assert "99.0% minimum" in output
    assert "the best was" in output


def test_the_report_says_where_the_workbook_went(workspace, capsys):
    run(workspace)
    assert "Saved to:" in capsys.readouterr().out


def test_a_rationale_is_grounded_in_the_real_figures():
    from conftest import make_record

    record = make_record(probability=0.841, odds=1.44, value=0.060)
    record.lambda_total = 3.30
    record.implied_fair = 0.781
    record.margin_adjusted = True
    record.home_form_2plus, record.home_form_matches = 9, 10
    record.away_form_2plus, record.away_form_matches = 9, 10
    record.h2h_matches, record.h2h_2plus = 10, 8

    text = build_rationale(record)
    assert "84.1%" in text
    assert "3.30" in text
    assert "9 of their last 10" in text
    assert "10 meetings went Over 1.5 8 times" in text
    assert "78.1%" in text and "+6.0" in text


def test_a_rationale_omits_statistics_that_are_missing():
    """Spec 12 - an absent statistic is left out, not described with a placeholder."""
    from conftest import make_record

    record = make_record(probability=0.80)
    record.h2h_matches = None
    record.home_form_2plus = None
    record.away_form_2plus = None

    text = build_rationale(record)
    assert "meetings" not in text
    assert "None" not in text


def test_the_report_closes_with_the_disclaimer(workspace, capsys):
    from config import DISCLAIMER

    run(workspace)
    assert DISCLAIMER in capsys.readouterr().out


def test_the_report_handles_a_run_that_found_nothing():
    from config import Config
    from models import AnalysisWindow
    from ranking import RankingSummary

    lines = build_report(
        RunOutcome(
            match_day=MATCH_DAY,
            window=AnalysisWindow.full_day(MATCH_DAY),
            config=Config(),
            fixtures_found=0,
            in_window=0,
            analysed=0,
            summary=RankingSummary(),
            records=[],
        )
    )
    text = "\n".join(lines)
    assert "NO MATCHES QUALIFIED." in text
    assert "No fixtures were collected at all" in text
